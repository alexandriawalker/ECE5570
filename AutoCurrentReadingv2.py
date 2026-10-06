import os
import time
from datetime import datetime

import serial  # pyserial
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

# ----------------------------- SETTINGS ------------------------------------
PORT = "COM4"
BAUD = 9600  # U1282A default. Must match the meter's setup menu.
INTERVAL_S = 60  # seconds between measurements
RUN_HOURS = 10  # script stops automatically after this many hours
NUM_AVG = 3  # readings averaged per measurement (matches the ESP8266 code)
AVG_GAP_S = 0.5  # pause between the readings so the meter has time to update
READ_CMD = "FETC?"  # returns what the meter is currently measuring.
# (Alternative: "READ?" triggers a fresh measurement.)
OUTPUT_FILE = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    f"dmm_log_discharging.xlsx",
)


# ---------------------------------------------------------------------------


def open_meter():
    """Open the serial port and confirm the meter responds."""
    ser = serial.Serial(
        port=PORT,
        baudrate=BAUD,
        bytesize=serial.EIGHTBITS,
        parity=serial.PARITY_NONE,
        stopbits=serial.STOPBITS_ONE,
        timeout=3,
    )
    time.sleep(0.5)
    ser.reset_input_buffer()
    idn = query(ser, "*IDN?")
    print(f"Connected to: {idn}")
    return ser


def query(ser, cmd):
    """Send a SCPI command and return the response line as text."""
    ser.reset_input_buffer()
    ser.write((cmd + "\n").encode("ascii"))
    reply = ser.readline().decode("ascii", errors="replace").strip()
    if not reply:
        raise TimeoutError(f"No response to '{cmd}'")
    return reply


def read_average(ser):
    """Take NUM_AVG readings and return their average.

    If some readings fail, the good ones are still averaged.
    If all of them fail, the last error is raised.
    """
    readings = []
    last_err = None
    for i in range(NUM_AVG):
        try:
            readings.append(float(query(ser, READ_CMD)))
        except (ValueError, TimeoutError) as err:
            last_err = err
        if i < NUM_AVG - 1:
            time.sleep(AVG_GAP_S)
    if not readings:
        raise last_err
    return sum(readings) / len(readings)


def create_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "DMM Log"
    ws.append(["Timestamp", "Reading"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 18
    return wb, ws


def save(wb):
    """Save, retrying if the file is open in Excel (which locks it)."""
    while True:
        try:
            wb.save(OUTPUT_FILE)
            return
        except PermissionError:
            print("  !! Close the Excel file so it can be saved. Retrying in 5 s...")
            time.sleep(5)


def main():
    wb, ws = create_workbook()
    ser = open_meter()
    print(f"Logging every {INTERVAL_S} s (average of {NUM_AVG} readings) to: {OUTPUT_FILE}")
    print(f"Will run for {RUN_HOURS} h. Press Ctrl+C (or the red Stop button) to stop early.\n")

    start = time.monotonic()
    end_time = start + RUN_HOURS * 3600
    next_time = start
    try:
        while time.monotonic() < end_time:
            stamp = datetime.now()
            try:
                reading = read_average(ser)
            except (ValueError, TimeoutError) as err:
                print(f"{stamp:%H:%M:%S}  read error: {err}")
                reading = None

            row = ws.max_row + 1
            ws.cell(row=row, column=1, value=stamp).number_format = "hh:mm:ss"
            ws.cell(row=row, column=2, value=reading if reading is not None else "ERROR")
            save(wb)
            print(f"{stamp:%H:%M:%S}  {reading}")

            # Keep a fixed 60 s cadence (no drift from read/save time)
            next_time += INTERVAL_S
            if next_time >= end_time:
                break
            time.sleep(max(0, next_time - time.monotonic()))
        print(f"\nReached {RUN_HOURS} hour run time. Finished.")
    except KeyboardInterrupt:
        print("\nStopped by user.")
    finally:
        save(wb)
        ser.close()
        print(f"Data saved to {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
